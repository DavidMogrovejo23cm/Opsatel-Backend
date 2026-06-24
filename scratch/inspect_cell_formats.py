import openpyxl

wb = openpyxl.load_workbook(r"../BASE DE DATOS 2026(3).xlsx", data_only=False)
sheet = wb.active

headers = [cell.value for cell in sheet[1]]
target_cols = ["PENDIENTE", "IPTV", "adicional", "total", "plan", "INTERNET PAY"]
col_indices = {name: headers.index(name) + 1 for name in target_cols if name in headers}

print("Checking formats of non-empty cells in relevant columns:")
for name, col_idx in col_indices.items():
    non_empty = []
    for r in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=r, column=col_idx)
        if cell.value is not None:
            non_empty.append((r, cell.value, cell.data_type, cell.number_format))
    print(f"\nColumn '{name}' has {len(non_empty)} non-empty cells out of {sheet.max_row - 1} total rows.")
    if non_empty:
        print("First 10 non-empty cells:")
        for item in non_empty[:10]:
            print(f"  Row {item[0]}: value={repr(item[1])}, excel_type={item[2]}, format={item[3]}")
