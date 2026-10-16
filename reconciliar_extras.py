import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, date
import os
import sys

def run():
    print("Iniciando script de reconciliación de Extras General...")
    
    antigua_path = "2026 ADMNISTRACION(6).xlsx"
    if not os.path.exists(antigua_path) and os.path.exists(os.path.join("..", antigua_path)):
        antigua_path = os.path.join("..", antigua_path)

    nueva_path = "Base_Datos_Extras_General_2026-10-15_backup_original.xlsx"
    if not os.path.exists(nueva_path) and os.path.exists(os.path.join("..", nueva_path)):
        nueva_path = os.path.join("..", nueva_path)
    elif not os.path.exists(nueva_path):
        nueva_path = "Base_Datos_Extras_General_2026-10-15.xlsx"
        if not os.path.exists(nueva_path) and os.path.exists(os.path.join("..", nueva_path)):
            nueva_path = os.path.join("..", nueva_path)
        
    wb_a = openpyxl.load_workbook(antigua_path, data_only=True)
    ws_a = wb_a['PLATAFORMA']

    wb_n = openpyxl.load_workbook(nueva_path, data_only=True)
    ws_n = wb_n['PLATAFORMA']

    def clean_num(v):
        if v is None: return 0.0
        if isinstance(v, (int, float)): return round(float(v), 2)
        s = str(v).strip().replace('$', '').replace(',', '.')
        try: return round(float(s), 2)
        except: return 0.0

    def fmt_date(v):
        if v is None: return None
        if hasattr(v, 'strftime'): return v.strftime('%d/%m/%Y')
        s = str(v).strip()
        if not s or s.upper() == 'NONE': return None
        return s

    def clean_str(v):
        if v is None: return None
        s = str(v).strip()
        if not s or s.upper() == 'NONE': return None
        return s

    # 1. Capturar pagos reales de OCTUBRE y NOVIEMBRE de la base exportada (del sistema)
    oct_payments = {}
    nov_payments = {}
    
    for r in range(2, ws_n.max_row + 1):
        nom = str(ws_n.cell(r, 2).value or '').strip()
        if not nom: continue
        nom_u = nom.upper()
        
        # Octubre en nueva es col 85 a 92 (pago es 89, fecha 88, banco 90, ref 91, fac 86)
        pago_oct = clean_num(ws_n.cell(r, 89).value)
        fecha_oct = fmt_date(ws_n.cell(r, 88).value)
        banco_oct = clean_str(ws_n.cell(r, 90).value)
        ref_oct = clean_str(ws_n.cell(r, 91).value)
        fac_oct = clean_str(ws_n.cell(r, 86).value)
        
        # Filtrar solo pagos legítimos de octubre (no los desplazados que decían IMPORTACION JUNIO/AGOSTO/MAYO)
        if pago_oct > 0 and (
            (fecha_oct and '10/2026' in str(fecha_oct)) or 
            (ref_oct and 'OCTUBRE' in ref_oct.upper()) or 
            (ref_oct and '2027' in ref_oct) or
            ('LIMA JUMA LUIS MIGUEL' in nom_u) or
            ('ARPI' in nom_u)
        ):
            oct_payments[nom_u] = {
                'pago': pago_oct,
                'fecha_pago': fecha_oct or '12/10/2026',
                'banco': banco_oct or 'PICHINCHA',
                'ref': ref_oct or 'Pago EXTRA - OCTUBRE',
                'factura': fac_oct
            }
            print(f"-> Pago Octubre detectado para {nom_u}: ${pago_oct} ({banco_oct}, {fecha_oct})")

        # Noviembre en nueva es col 93 a 100 (pago es 97, fecha 96, banco 98, ref 99)
        pago_nov = clean_num(ws_n.cell(r, 97).value)
        fecha_nov = fmt_date(ws_n.cell(r, 96).value)
        banco_nov = clean_str(ws_n.cell(r, 98).value)
        ref_nov = clean_str(ws_n.cell(r, 99).value)
        if pago_nov > 0 and (fecha_nov and '10/2026' in str(fecha_nov)):
            nov_payments[nom_u] = {
                'pago': pago_nov,
                'fecha_pago': fecha_nov,
                'banco': banco_nov or 'PICHINCHA',
                'ref': ref_nov or 'Pago EXTRA - NOVIEMBRE',
                'factura': clean_str(ws_n.cell(r, 94).value)
            }
            print(f"-> Pago Noviembre detectado para {nom_u}: ${pago_nov} ({banco_nov}, {fecha_nov})")

    # Meses y su especificación en Antigua
    # Enero..Junio: 8 columnas (base, factura, fecha_a_pagar, fecha_pago, pago, banco, cod, saldo)
    # Julio..Octubre: 7 columnas (base, col_fac, col_f_pago, col_pago, col_banco, col_cod, col_saldo)
    meses_spec = [
        ('ENERO', 13, 8),
        ('FEBRERO', 21, 8),
        ('MARZO', 29, 8),
        ('ABRIL', 37, 8),
        ('MAYO', 45, 8),
        ('JUNIO', 53, 8),
        ('JULIO', 61, 7),
        ('AGOSTO', 68, 7),
        ('SEPTIEMBRE', 75, 7),
        ('OCTUBRE', 82, 7)
    ]

    all_clients_data = []

    for r in range(2, ws_a.max_row + 1):
        cod = clean_str(ws_a.cell(r, 1).value)
        nom = clean_str(ws_a.cell(r, 2).value)
        if not nom and not cod:
            continue
        
        # Ignorar fila si no tiene nombre
        if not nom:
            continue
            
        contacto = clean_str(ws_a.cell(r, 3).value)
        proveedor = clean_str(ws_a.cell(r, 4).value) or "OPSATEL"
        usuario = clean_str(ws_a.cell(r, 5).value)
        contrasena = clean_str(ws_a.cell(r, 6).value) or "TV2026.@"
        cuentas = clean_str(ws_a.cell(r, 7).value) or "1"
        mac = clean_str(ws_a.cell(r, 8).value)
        obs = clean_str(ws_a.cell(r, 9).value)
        estado = clean_str(ws_a.cell(r, 10).value) or "FIJO"
        valor = clean_num(ws_a.cell(r, 11).value)
        if valor == 0.0:
            c11_raw = str(ws_a.cell(r, 11).value or '').strip().upper()
            if c11_raw != 'GRATIS' and 'TECNICO' not in nom.upper():
                m_fees = [clean_num(ws_a.cell(r, c).value) for c in [13, 21, 29, 37, 45, 53, 61, 68, 75] if clean_num(ws_a.cell(r, c).value) > 0]
                if m_fees:
                    valor = max(set(m_fees), key=m_fees.count)
        activo = clean_str(ws_a.cell(r, 12).value) or "SI"

        nom_u = nom.upper()
        
        client_record = {
            'cod': cod,
            'nombre_cliente': nom,
            'contacto': contacto,
            'proveedor': proveedor,
            'usuario': usuario,
            'contrasena': contrasena,
            'cuentas': cuentas,
            'mac_smart_one': mac,
            'observaciones': obs,
            'estado': estado,
            'valor': valor,
            'activo': activo,
            'meses': {}
        }

        # Extraer Enero a Septiembre
        for m_name, c_start, c_len in meses_spec:
            if m_name == 'OCTUBRE':
                continue # Procesar Octubre abajo con las reglas especiales
                
            fee = clean_num(ws_a.cell(r, c_start).value)
            
            if c_len == 8:
                factura = clean_str(ws_a.cell(r, c_start + 1).value)
                f_a_pagar = fmt_date(ws_a.cell(r, c_start + 2).value)
                f_pago = fmt_date(ws_a.cell(r, c_start + 3).value)
                pago = clean_num(ws_a.cell(r, c_start + 4).value)
                banco = clean_str(ws_a.cell(r, c_start + 5).value)
                cod_ref = clean_str(ws_a.cell(r, c_start + 6).value)
                saldo = clean_num(ws_a.cell(r, c_start + 7).value)
            else:
                # 7 columnas (JULIO, AGOSTO, SEPTIEMBRE)
                col2 = ws_a.cell(r, c_start + 1).value # Factura o fecha a pagar
                col3 = ws_a.cell(r, c_start + 2).value # Fecha a pagar o fecha pago
                col4 = ws_a.cell(r, c_start + 3).value # PAGO
                col5 = ws_a.cell(r, c_start + 4).value # BANCO
                col6 = ws_a.cell(r, c_start + 5).value # COD
                col7 = ws_a.cell(r, c_start + 6).value # SALDO
                
                factura = None
                f_a_pagar = None
                f_pago = None
                
                # Determinar si col2 es factura (texto con letras como C02958, 001-...) o fecha
                if col2 is not None:
                    s_col2 = str(col2).strip()
                    if hasattr(col2, 'strftime') or (len(s_col2) >= 8 and ('/' in s_col2 or '-' in s_col2) and not s_col2.startswith('001') and not s_col2.startswith('C')):
                        f_a_pagar = fmt_date(col2)
                    else:
                        factura = s_col2
                        
                if col3 is not None:
                    f_pago = fmt_date(col3)
                elif f_a_pagar and not factura:
                    f_pago = f_a_pagar
                    f_a_pagar = None
                    
                pago = clean_num(col4)
                banco = clean_str(col5)
                cod_ref = clean_str(col6)
                saldo = clean_num(col7)

            client_record['meses'][m_name] = {
                'fee': fee,
                'factura': factura,
                'fecha_a_pagar': f_a_pagar,
                'fecha_pago': f_pago,
                'pago': pago,
                'banco': banco,
                'cod': cod_ref,
                'saldo': saldo
            }

        # === PROCESAR OCTUBRE ===
        # Saldo del mes anterior (Septiembre)
        sep_saldo = client_record['meses']['SEPTIEMBRE']['saldo']
        
        # Tarifa de Octubre: de col CD(82) si existe, o Septiembre fee, o valor base
        oct_fee = clean_num(ws_a.cell(r, 82).value)
        if oct_fee == 0.0:
            oct_fee = client_record['meses']['SEPTIEMBRE']['fee']
        if oct_fee == 0.0 and estado != 'ELIMINADO':
            oct_fee = valor
            
        if nom_u in oct_payments:
            oct_info = oct_payments[nom_u]
            oct_pago = oct_info['pago']
            oct_f_pago = oct_info['fecha_pago']
            oct_banco = oct_info['banco']
            oct_ref = oct_info['ref']
            oct_fac = oct_info['factura']
            # Running saldo: septiembre_saldo + fee - pago
            oct_saldo = round(sep_saldo + oct_fee - oct_pago, 2)
        else:
            oct_pago = 0.0
            oct_f_pago = None
            oct_banco = None
            oct_ref = None
            oct_fac = None
            # Running saldo: septiembre_saldo + fee
            oct_saldo = round(sep_saldo + oct_fee, 2)

        client_record['meses']['OCTUBRE'] = {
            'fee': oct_fee,
            'factura': oct_fac,
            'fecha_a_pagar': None,
            'fecha_pago': oct_f_pago,
            'pago': oct_pago,
            'banco': oct_banco,
            'cod': oct_ref,
            'saldo': oct_saldo
        }

        # === PROCESAR NOVIEMBRE ===
        nov_fee = oct_fee if estado != 'ELIMINADO' else 0.0
        if nom_u in nov_payments:
            nov_info = nov_payments[nom_u]
            nov_pago = nov_info['pago']
            nov_f_pago = nov_info['fecha_pago']
            nov_banco = nov_info['banco']
            nov_ref = nov_info['ref']
            nov_fac = nov_info['factura']
            nov_saldo = round(oct_saldo + nov_fee - nov_pago, 2)
        else:
            nov_pago = 0.0
            nov_f_pago = None
            nov_banco = None
            nov_ref = None
            nov_fac = None
            nov_saldo = round(oct_saldo + nov_fee, 2)

        client_record['meses']['NOVIEMBRE'] = {
            'fee': nov_fee,
            'factura': nov_fac,
            'fecha_a_pagar': None,
            'fecha_pago': nov_f_pago,
            'pago': nov_pago,
            'banco': nov_banco,
            'cod': nov_ref,
            'saldo': nov_saldo
        }

        # === PROCESAR DICIEMBRE ===
        dic_fee = oct_fee if estado != 'ELIMINADO' else 0.0
        dic_saldo = round(nov_saldo + dic_fee, 2)
        client_record['meses']['DICIEMBRE'] = {
            'fee': dic_fee,
            'factura': None,
            'fecha_a_pagar': None,
            'fecha_pago': None,
            'pago': 0.0,
            'banco': None,
            'cod': None,
            'saldo': dic_saldo
        }

        all_clients_data.append(client_record)

    print(f"Total clientes procesados y reconciliados: {len(all_clients_data)}")

    # 2. Generar el Excel definitivo Base_Datos_Extras_General_2026-10-15.xlsx
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "PLATAFORMA"

    meses_nombres = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
    
    headers = [
        "COD", "NOMBRE CLIENTE", "CONTACTO", "PROVEEDOR", "USUARIO", "CONTRASEÑA",
        "CUENTAS", "MAC SMART ONE", "OBSERVACIONES", "ESTADO", "VALOR", "ACTIVO"
    ]
    for m in meses_nombres:
        headers.extend([m, "FACTURA", "FECHA A PAGAR", "FECHA DE PAGO", "PAGO", "BANCO", "COD", "SALDO"])

    ws_out.append(headers)

    # Estilos del encabezado
    header_fill_general = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_fill_month = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for c in range(1, len(headers) + 1):
        cell = ws_out.cell(1, c)
        cell.font = header_font
        cell.alignment = align_center
        if c <= 12:
            cell.fill = header_fill_general
        else:
            cell.fill = header_fill_month

    for client in all_clients_data:
        row_vals = [
            client['cod'],
            client['nombre_cliente'],
            client['contacto'],
            client['proveedor'],
            client['usuario'],
            client['contrasena'],
            client['cuentas'],
            client['mac_smart_one'],
            client['observaciones'],
            client['estado'],
            client['valor'],
            client['activo']
        ]
        for m in meses_nombres:
            m_data = client['meses'][m]
            row_vals.extend([
                m_data['fee'],
                m_data['factura'],
                m_data['fecha_a_pagar'],
                m_data['fecha_pago'],
                m_data['pago'],
                m_data['banco'],
                m_data['cod'],
                m_data['saldo']
            ])
        ws_out.append(row_vals)

    # Formatear filas de datos
    for r in range(2, ws_out.max_row + 1):
        for c in range(1, len(headers) + 1):
            cell = ws_out.cell(r, c)
            cell.border = thin_border
            h = headers[c - 1]
            if h in ["VALOR", "PAGO", "SALDO"] or h in meses_nombres:
                cell.alignment = align_right
                cell.number_format = '#,##0.00;[Red]-#,##0.00;0.00'
                if h == "SALDO" and isinstance(cell.value, (int, float)):
                    if cell.value < 0:
                        cell.font = Font(name="Calibri", size=10, bold=True, color="166534") # Verde oscuro para excedente
                        cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                    elif cell.value > 0:
                        cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B") # Rojo oscuro para deuda
                        cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                elif h == "PAGO" and isinstance(cell.value, (int, float)) and cell.value > 0:
                    cell.font = Font(name="Calibri", size=10, bold=True, color="047857")
            elif h in ["COD", "CUENTAS", "ESTADO", "ACTIVO", "FECHA A PAGAR", "FECHA DE PAGO"]:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

    # Congelar paneles en columnas A y B (scroll horizontal cómodo)
    ws_out.freeze_panes = "C2"

    # Autoajuste de anchos de columna
    for col in ws_out.columns:
        col_letter = col[0].column_letter
        c_idx = col[0].column
        h = headers[c_idx - 1]
        max_len = max(len(str(cell.value or '')) for cell in col)
        if c_idx == 2: # Nombre
            ws_out.column_dimensions[col_letter].width = 30
        elif c_idx == 9: # Observaciones
            ws_out.column_dimensions[col_letter].width = 25
        elif h in ["PAGO", "SALDO", "VALOR"] or h in meses_nombres:
            ws_out.column_dimensions[col_letter].width = 12
        elif "FECHA" in h:
            ws_out.column_dimensions[col_letter].width = 13
        elif h == "BANCO":
            ws_out.column_dimensions[col_letter].width = 14
        else:
            ws_out.column_dimensions[col_letter].width = max(max_len + 3, 10)

    # Guardar sobre Base_Datos_Extras_General_2026-10-15.xlsx
    target_excel = "Base_Datos_Extras_General_2026-10-15.xlsx"
    if not os.path.exists(target_excel) and os.path.exists(os.path.join("..", target_excel)):
        target_excel = os.path.join("..", target_excel)
    wb_out.save(target_excel)
    print(f"Excel guardado exitosamente en: {target_excel}")

    # 3. Sincronizar directamente con la base de datos MySQL (y SQLite si aplica)
    try:
        sys.path.append('Opsatel-Backend')
        import database
        import models
        from sqlalchemy.orm import Session

        db = database.SessionLocal()
        print("Sincronizando con base de datos...")

        # Limpiar y re-poblar extras
        db.query(models.PagoExtra).delete()
        db.query(models.ClienteExtra).delete()
        db.commit()

        for c_data in all_clients_data:
            cliente = models.ClienteExtra(
                cod=c_data['cod'],
                nombre_cliente=c_data['nombre_cliente'],
                contacto=c_data['contacto'],
                proveedor=c_data['proveedor'],
                usuario=c_data['usuario'],
                contrasena=c_data['contrasena'],
                cuentas=c_data['cuentas'],
                mac_smart_one=c_data['mac_smart_one'],
                observaciones=c_data['observaciones'],
                estado=c_data['estado'],
                valor=c_data['valor'],
                activo=c_data['activo'],
                fecha_ingreso="2026-01-01"
            )
            
            total_pagado = 0.0
            for m in meses_nombres:
                m_low = m.lower()
                m_info = c_data['meses'][m]
                setattr(cliente, f"{m_low}_factura", m_info['factura'])
                setattr(cliente, f"{m_low}_fecha_a_pagar", m_info['fecha_a_pagar'])
                setattr(cliente, f"{m_low}_fecha_pago", m_info['fecha_pago'])
                setattr(cliente, f"{m_low}_pago", m_info['pago'])
                setattr(cliente, f"{m_low}_banco", m_info['banco'])
                setattr(cliente, f"{m_low}_cod", m_info['cod'])
                setattr(cliente, f"{m_low}_saldo", m_info['saldo'])
                total_pagado += m_info['pago']
                
            cliente.total_pagado = round(total_pagado, 2)
            # Saldo pendiente actual (Octubre)
            oct_s = c_data['meses']['OCTUBRE']['saldo']
            cliente.saldo_pendiente = max(0.0, oct_s)

            db.add(cliente)
            db.flush()

            # Registrar en historial_pagos_extras
            for m in meses_nombres:
                m_info = c_data['meses'][m]
                if m_info['pago'] > 0:
                    f_dt = datetime(2026, meses_nombres.index(m) + 1, 15)
                    if m_info['fecha_pago']:
                        try:
                            parts = str(m_info['fecha_pago']).split('/')
                            if len(parts) == 3:
                                f_dt = datetime(int(parts[2]), int(parts[1]), int(parts[0]))
                        except:
                            pass
                    pago_rec = models.PagoExtra(
                        cliente_id=cliente.id,
                        monto=m_info['pago'],
                        metodo_pago=m_info['banco'] or 'EFECTIVO',
                        mes_correspondiente=m,
                        referencia=m_info['cod'] or f"Pago EXTRA - {m}",
                        factura=m_info['factura'],
                        estado="Completado",
                        fecha_pago=f_dt,
                        anulado=False
                    )
                    db.add(pago_rec)

        db.commit()
        db.close()
        print("Base de datos sincronizada con éxito total!")
    except Exception as db_err:
        print(f"Aviso BD: {db_err}")

    print("Reconciliación completada exitosamente.")

if __name__ == '__main__':
    run()
