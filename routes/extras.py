# pyrefly: ignore [missing-import]
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File, Form
from fastapi.responses import StreamingResponse
# pyrefly: ignore [missing-import]
from sqlalchemy.orm import Session
from typing import List
import models, schemas
from database import get_db
from datetime import datetime, date
import asyncio
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import unicodedata
import re
from services.xui_service import create_xui_user, delete_xui_user
from .auth import get_current_user, require_role

router = APIRouter(prefix="/extras", tags=["extras"])


@router.get("/", response_model=List[schemas.ClienteExtraResponse])
def listar_extras(db: Session = Depends(get_db)):
    extras = db.query(models.ClienteExtra).all()
    for e in extras:
        if str(e.estado or "").strip().upper() == "ELIMINADO":
            e.saldo_pendiente = 0.0
            for m in ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]:
                setattr(e, f"{m}_saldo", 0.0)
    return extras

@router.post("/", response_model=schemas.ClienteExtraResponse, dependencies=[Depends(require_role(["administrador", "secretario", "tecnico"]))])
async def crear_extra(extra: schemas.ClienteExtraCreate, db: Session = Depends(get_db)):
    if not extra.fecha_ingreso:
        extra.fecha_ingreso = datetime.now().strftime("%Y-%m-%d")
    db_extra = models.ClienteExtra(**extra.dict())
    db.add(db_extra)
    db.commit()
    db.refresh(db_extra)

    # Recalcular cuadrícula inicial para este nuevo cliente extra
    recalcular_cuadricula_extra(db_extra, db)
    db.commit()
    db.refresh(db_extra)

    # Si posee credenciales para IPTV, aprovisionar en el panel XUI de forma asíncrona
    if db_extra.usuario and str(db_extra.usuario).strip() and db_extra.contrasena:
        iptv_u = str(db_extra.usuario).strip()
        iptv_p = str(db_extra.contrasena).strip()
        try:
            cuentas = int(db_extra.cuentas or 1)
        except (ValueError, TypeError):
            cuentas = 1

        try:
            res = await create_xui_user(
                username=iptv_u,
                password=iptv_p,
                max_connections=cuentas,
                bouquets=["1", "2", "5"],
                allowed_outputs=["1", "2"]
            )
            print(f"IPTV EXTRA: Cuenta de XUI creada exitosamente para {iptv_u}. Detalle: {res}")
        except Exception as xui_err:
            print(f"ERROR IPTV EXTRA: Falló la creación en panel XUI para {iptv_u}: {type(xui_err).__name__}: {str(xui_err) or repr(xui_err)}")

    return db_extra

@router.patch("/{id}", response_model=schemas.ClienteExtraResponse, dependencies=[Depends(require_role(["administrador", "secretario", "tecnico"]))])
def actualizar_extra(id: int, data: schemas.ClienteExtraUpdate, db: Session = Depends(get_db)):
    db_extra = db.query(models.ClienteExtra).filter(models.ClienteExtra.id == id).first()
    if not db_extra:
        raise HTTPException(status_code=404, detail="Cliente extra no encontrado")
    
    update_data = data.dict(exclude_unset=True)
    meses_validos = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio", 
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"
    ]

    for var, value in update_data.items():
        if var == "valor" or var.endswith("_pago") or var.endswith("_saldo") or var in ("saldo_pendiente", "total_pagado"):
            try:
                value = float(value or 0.0)
            except (ValueError, TypeError):
                value = 0.0
        elif var.endswith("_fecha_pago") or var == "fecha_ingreso":
            if value and str(value).strip():
                value = str(value).strip()
        elif var == "activo":
            value = str(value or "SI").strip().upper()
        elif var == "estado":
            value = str(value or "30 DÍAS").strip().upper()
        setattr(db_extra, var, value)

    # Identificar meses afectados por la edición
    meses_tocados = set()
    for var in update_data.keys():
        for m in meses_validos:
            if var.startswith(f"{m}_"):
                meses_tocados.add(m)
                break

    # Sincronizar PagoExtra para que el historial y balance reflejen los pagos según el mes
    for m in meses_tocados:
        if f"{m}_pago" in update_data and f"{m}_saldo" not in update_data:
            v_base = float(db_extra.valor or 0.0)
            p_val = float(getattr(db_extra, f"{m}_pago", 0.0) or 0.0)
            setattr(db_extra, f"{m}_saldo", round(v_base - p_val, 2))

        m_pago = float(getattr(db_extra, f"{m}_pago", 0.0) or 0.0)
        m_banco = getattr(db_extra, f"{m}_banco", None) or "EFECTIVO"
        m_factura = getattr(db_extra, f"{m}_factura", None)
        m_cod = getattr(db_extra, f"{m}_cod", None)
        m_fecha = getattr(db_extra, f"{m}_fecha_pago", None)

        pago_existente = db.query(models.PagoExtra).filter(
            models.PagoExtra.cliente_id == db_extra.id,
            models.PagoExtra.mes_correspondiente == m.upper(),
            models.PagoExtra.anulado == False
        ).first()

        if m_pago > 0:
            f_pago_dt = None
            if m_fecha:
                for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
                    try:
                        f_pago_dt = datetime.strptime(str(m_fecha).strip(), fmt)
                        break
                    except Exception:
                        pass
            if not f_pago_dt:
                m_idx = meses_validos.index(m) + 1
                anio_actual = datetime.now().year
                f_pago_dt = datetime(anio_actual, m_idx, 15)

            if pago_existente:
                pago_existente.monto = m_pago
                pago_existente.metodo_pago = m_banco
                pago_existente.factura = m_factura
                pago_existente.referencia = m_cod or f"PAGO EXTRA {m.upper()}"
                pago_existente.fecha_pago = f_pago_dt
                pago_existente.estado = "Completado"
            else:
                nuevo_pago = models.PagoExtra(
                    cliente_id=db_extra.id,
                    monto=m_pago,
                    metodo_pago=m_banco,
                    mes_correspondiente=m.upper(),
                    referencia=m_cod or f"PAGO EXTRA {m.upper()}",
                    factura=m_factura,
                    estado="Completado",
                    fecha_pago=f_pago_dt,
                    anulado=False
                )
                db.add(nuevo_pago)
        else:
            if pago_existente:
                db.delete(pago_existente)

    is_eliminado = str(db_extra.estado or "").strip().upper() == "ELIMINADO"
    if is_eliminado:
        db_extra.saldo_pendiente = 0.0
        for m in meses_validos:
            setattr(db_extra, f"{m}_saldo", 0.0)
    else:
        recalcular_cuadricula_extra(db_extra, db)

    # Actualizar totales globales
    total_pagado = 0.0
    for m in meses_validos:
        total_pagado += float(getattr(db_extra, f"{m}_pago", 0.0) or 0.0)
    db_extra.total_pagado = round(total_pagado, 2)

    if is_eliminado:
        db_extra.saldo_pendiente = 0.0
    elif "saldo_pendiente" in update_data:
        db_extra.saldo_pendiente = float(update_data["saldo_pendiente"] or 0.0)
    elif "saldo" in update_data:
        db_extra.saldo_pendiente = float(update_data["saldo"] or 0.0)
    else:
        mes_actual_idx = datetime.now().month - 1
        saldo_mes_actual = float(getattr(db_extra, f"{meses_validos[mes_actual_idx]}_saldo", 0.0) or 0.0)
        db_extra.saldo_pendiente = max(0.0, saldo_mes_actual)

    db.commit()
    db.refresh(db_extra)
    return db_extra

@router.delete("/all", dependencies=[Depends(require_role(["administrador"]))])
def eliminar_todos_extras(db: Session = Depends(get_db)):
    """
    Elimina TODOS los clientes extras de la base de datos y sus historiales de pagos.
    Solo accesible para administradores.
    """
    try:
        db.query(models.PagoExtra).delete()
        count = db.query(models.ClienteExtra).delete()
        db.commit()

        try:
            from sqlalchemy import text
            from database import engine
            is_postgresql = "postgresql" in str(engine.url).lower() or "psycopg" in str(engine.url).lower()
            is_mysql = "mysql" in str(engine.url).lower()
            if is_postgresql:
                db.execute(text("ALTER SEQUENCE clientes_extras_id_seq RESTART WITH 1"))
                db.execute(text("ALTER SEQUENCE historial_pagos_extras_id_seq RESTART WITH 1"))
            elif is_mysql:
                db.execute(text("ALTER TABLE clientes_extras AUTO_INCREMENT = 1"))
                db.execute(text("ALTER TABLE historial_pagos_extras AUTO_INCREMENT = 1"))
            db.commit()
        except Exception:
            pass

        return {"message": f"Todos los clientes extras ({count}) han sido eliminados correctamente."}
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error al eliminar clientes extras: {str(e)}")

@router.delete("/{id}", dependencies=[Depends(require_role(["administrador"]))])
async def eliminar_extra(id: int, db: Session = Depends(get_db)):
    db_extra = db.query(models.ClienteExtra).filter(models.ClienteExtra.id == id).first()
    if not db_extra:
        raise HTTPException(status_code=404, detail="Cliente extra no encontrado")
    
    usuario_xui = db_extra.usuario
    db.delete(db_extra)
    db.commit()

    if usuario_xui and str(usuario_xui).strip():
        try:
            res = await delete_xui_user(str(usuario_xui).strip())
            print(f"IPTV EXTRA: Eliminación en XUI para {usuario_xui}: {res}")
        except Exception as xui_err:
            print(f"ERROR IPTV EXTRA: Falló eliminación en XUI para {usuario_xui}: {xui_err}")

    return {"message": "Cliente extra eliminado"}

def recalcular_cuadricula_extra(cliente: models.ClienteExtra, db: Session):
    meses_validos = [
        "enero", "febrero", "marzo", "abril", "mayo", "junio", 
        "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"
    ]
    
    is_eliminado = str(cliente.estado or "").strip().upper() == "ELIMINADO"
    if is_eliminado:
        cliente.saldo_pendiente = 0.0
        for m in meses_validos:
            setattr(cliente, f"{m}_saldo", 0.0)
        return

    valor_mensual = float(cliente.valor or 0)
    
    limite_ingreso_idx = 0
    if cliente.fecha_ingreso:
        try:
            mes_ingreso_num = int(cliente.fecha_ingreso.split("-")[1])
            limite_ingreso_idx = mes_ingreso_num - 1
        except:
            limite_ingreso_idx = 0

    # 1. Cargar pagos válidos
    pagos = db.query(models.PagoExtra).filter(
        models.PagoExtra.cliente_id == cliente.id,
        models.PagoExtra.anulado == False,
        models.PagoExtra.estado == "Completado"
    ).order_by(models.PagoExtra.fecha_pago.asc(), models.PagoExtra.id.asc()).all()
    
    # Agrupar pagos por mes correspondiente (PER MES)
    pagos_por_mes = {}
    for p in pagos:
        mes_key = (p.mes_correspondiente or "").strip().lower()
        if mes_key in meses_validos:
            if mes_key not in pagos_por_mes:
                pagos_por_mes[mes_key] = []
            pagos_por_mes[mes_key].append(p)

    running_saldo = 0.0
    for i in range(12):
        mes_actual = meses_validos[i]
        if i < limite_ingreso_idx:
            # Mes anterior al ingreso del cliente
            setattr(cliente, f"{mes_actual}_saldo", 0.0)
            setattr(cliente, f"{mes_actual}_pago", 0.0)
        else:
            if mes_actual in pagos_por_mes:
                total_mes = sum(float(p.monto or 0) for p in pagos_por_mes[mes_actual])
                ultimo_pago = pagos_por_mes[mes_actual][-1]
                setattr(cliente, f"{mes_actual}_pago", round(total_mes, 2))
                setattr(cliente, f"{mes_actual}_fecha_pago", ultimo_pago.fecha_pago.strftime("%d/%m/%Y") if ultimo_pago.fecha_pago else datetime.now().strftime("%d/%m/%Y"))
                setattr(cliente, f"{mes_actual}_banco", ultimo_pago.metodo_pago)
                setattr(cliente, f"{mes_actual}_factura", ultimo_pago.factura)
                setattr(cliente, f"{mes_actual}_cod", ultimo_pago.referencia)
            else:
                total_mes = float(getattr(cliente, f"{mes_actual}_pago", 0) or 0)

            # Cuota mensual del cliente activo
            fee_mes = valor_mensual if str(cliente.activo).strip().upper() != 'NO' and not is_eliminado else 0.0
            # Running balance con arrastre de excedentes:
            running_saldo = round(running_saldo + fee_mes - total_mes, 2)
            setattr(cliente, f"{mes_actual}_saldo", running_saldo)

    # Calcular totales globales
    total_pagado = 0.0
    for m in meses_validos:
        total_pagado += float(getattr(cliente, f"{m}_pago", 0) or 0)
    cliente.total_pagado = round(total_pagado, 2)
    # Saldo pendiente representa la deuda neta al mes actual
    mes_actual_idx = datetime.now().month - 1
    saldo_mes_actual = float(getattr(cliente, f"{meses_validos[mes_actual_idx]}_saldo", 0) or 0)
    cliente.saldo_pendiente = max(0.0, saldo_mes_actual) if not is_eliminado else 0.0

@router.post("/{id}/pagar")
def registrar_pago_extra(
    id: int, 
    pago_data: schemas.PagoExtraCreate, 
    db: Session = Depends(get_db), 
    current_user: models.Usuario = Depends(require_role(["administrador", "secretario"]))
):
    cliente = db.query(models.ClienteExtra).filter(models.ClienteExtra.id == id).first()
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente extra no encontrado")

    # Obtener turno de caja abierto (si existe)
    turno = db.query(models.TurnoCaja).filter(
        models.TurnoCaja.usuario_id == current_user.id,
        models.TurnoCaja.estado == "Abierto"
    ).first()
    turno_id = turno.id if turno else None

    monto = float(pago_data.monto)
    mes_str = (pago_data.mes_correspondiente or "OCTUBRE").strip().upper()

    # Registrar en historial
    nuevo_pago = models.PagoExtra(
        cliente_id=id,
        monto=monto,
        metodo_pago=pago_data.metodo_pago,
        mes_correspondiente=mes_str,
        referencia=pago_data.referencia,
        factura=pago_data.factura,
        estado=pago_data.estado or "Completado",
        turnocaja_id=turno_id
    )
    db.add(nuevo_pago)
    db.flush()

    if nuevo_pago.estado == "Pendiente_Verificacion":
        db.commit()
        return {"message": "Pago extra registrado y pendiente de verificación bancaria."}

    # Recalcular cuadrícula si es completado
    recalcular_cuadricula_extra(cliente, db)
    db.commit()
    return {"message": "Pago extra registrado correctamente"}

@router.post("/{id}/pagos/{pago_id}/confirmar")
def confirmar_pago_extra(
    id: int, 
    pago_id: int, 
    db: Session = Depends(get_db), 
    current_user: models.Usuario = Depends(require_role(["administrador", "secretario"]))
):
    pago = db.query(models.PagoExtra).filter(models.PagoExtra.id == pago_id, models.PagoExtra.cliente_id == id).first()
    if not pago:
        raise HTTPException(status_code=404, detail="Pago no encontrado")
    if pago.estado != "Pendiente_Verificacion":
        raise HTTPException(status_code=400, detail="El pago ya está confirmado o anulado.")
    if pago.anulado:
        raise HTTPException(status_code=400, detail="No se puede confirmar un pago anulado.")
        
    cliente = db.query(models.ClienteExtra).filter(models.ClienteExtra.id == id).first()
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente extra no encontrado")
        
    pago.estado = "Completado"
    recalcular_cuadricula_extra(cliente, db)
    db.commit()
    return {"message": "Pago extra confirmado y aplicado exitosamente."}

@router.post("/{id}/pagos/{pago_id}/anular")
def anular_pago_extra(
    id: int, 
    pago_id: int, 
    request_data: schemas.PagoAnularRequest, 
    db: Session = Depends(get_db), 
    current_user: models.Usuario = Depends(require_role(["administrador", "secretario"]))
):
    pago = db.query(models.PagoExtra).filter(models.PagoExtra.id == pago_id, models.PagoExtra.cliente_id == id).first()
    if not pago:
        raise HTTPException(status_code=404, detail="Pago no encontrado")
    if pago.anulado:
        raise HTTPException(status_code=400, detail="Este pago ya se encuentra anulado.")
        
    cliente = db.query(models.ClienteExtra).filter(models.ClienteExtra.id == id).first()
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente extra no encontrado")
        
    pago.anulado = True
    pago.fecha_anulacion = datetime.utcnow()
    pago.anulado_por = current_user.username
    pago.motivo_anulacion = request_data.motivo_anulacion
    
    recalcular_cuadricula_extra(cliente, db)
    db.commit()
    return {"message": "Pago extra anulado exitosamente."}

@router.get("/pagos/historial")
def historial_pagos_extras(db: Session = Depends(get_db)):
    return db.query(models.PagoExtra).order_by(models.PagoExtra.fecha_pago.desc()).all()


# ========================================================================
# IMPORTACIÓN Y EXPORTACIÓN DE EXTRAS (TOTALMENTE APARTADO DE CLIENTES NORMALES)
# ========================================================================

def _norm_str(text):
    if not text:
        return ""
    s = unicodedata.normalize('NFKD', str(text)).encode('ASCII', 'ignore').decode('utf-8')
    return re.sub(r'[^A-Z0-9]', '', s.upper())


@router.post("/upload-db", dependencies=[Depends(require_role(["administrador", "secretario"]))])
def upload_database_extras(
    file: UploadFile = File(...),
    modo: str = Form("merge"), # "merge" (actualizar o crear) o "replace" (limpiar antes)
    db: Session = Depends(get_db)
):
    """
    Importa la base de datos de Clientes Extras (Extra General / Plataforma).
    Completamente separada y aislada de la base de datos de clientes normal.
    Soporta formato Excel con cuadrícula mensual completa (Enero..Diciembre).
    """
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="El archivo debe ser un Excel (.xlsx, .xls)")

    try:
        wb = openpyxl.load_workbook(file.file, data_only=True)
        # Buscar hoja PLATAFORMA o tomar la primera hoja activa
        sheet = None
        for sname in wb.sheetnames:
            if "PLATAFORMA" in sname.upper():
                sheet = wb[sname]
                break
        if sheet is None:
            sheet = wb.active

        if modo == "replace":
            db.query(models.PagoExtra).delete()
            db.query(models.ClienteExtra).delete()
            db.commit()

        months_order = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']
        month_positions = []
        for c in range(1, sheet.max_column + 1):
            h = sheet.cell(1, c).value
            if h and _norm_str(h) in months_order:
                month_positions.append((_norm_str(h), c))

        def clean_val(v):
            if v is None: return None
            if isinstance(v, (datetime, date)):
                return v.strftime('%d/%m/%Y')
            s = str(v).strip()
            return s if s else None

        def clean_num(v):
            if v is None: return 0.0
            if isinstance(v, (int, float)): return round(float(v), 2)
            s = str(v).strip().replace('$', '').replace(',', '.')
            if s.upper() == 'GRATIS': return 0.0
            try: return round(float(s), 2)
            except: return 0.0

        count_nuevos = 0
        count_actualizados = 0

        existing_extras = db.query(models.ClienteExtra).all()
        by_cod = {e.cod.strip().upper(): e for e in existing_extras if e.cod}
        by_nombre = {_norm_str(e.nombre_cliente): e for e in existing_extras if e.nombre_cliente}

        for r in range(2, sheet.max_row + 1):
            cod_raw = clean_val(sheet.cell(r, 1).value)
            nom_raw = clean_val(sheet.cell(r, 2).value)

            if not cod_raw and not nom_raw:
                continue

            contacto_raw = clean_val(sheet.cell(r, 3).value)
            proveedor_raw = clean_val(sheet.cell(r, 4).value) or "OPSATEL"
            usuario_raw = clean_val(sheet.cell(r, 5).value)
            contrasena_raw = clean_val(sheet.cell(r, 6).value) or "TV2026.@"
            cuentas_raw = str(sheet.cell(r, 7).value or '1').strip()
            mac_raw = clean_val(sheet.cell(r, 8).value)
            obs_raw = clean_val(sheet.cell(r, 9).value)
            estado_raw = clean_val(sheet.cell(r, 10).value) or "FIJO"
            valor_raw = clean_num(sheet.cell(r, 11).value)
            activo_raw = clean_val(sheet.cell(r, 12).value) or "SI"

            # Parsear datos de cada mes
            meses_data = {}
            first_active_month_idx = None
            month_fees = []

            for i, (m_name, start_c) in enumerate(month_positions):
                end_c = month_positions[i+1][1] - 1 if i+1 < len(month_positions) else sheet.max_column
                has_banco_header = any(str(sheet.cell(1, c).value or '').strip().upper() == 'BANCO' for c in range(start_c, end_c + 1))

                m_info = {
                    'factura': None,
                    'fecha_a_pagar': None,
                    'fecha_pago': None,
                    'pago': 0.0,
                    'banco': None,
                    'cod': None,
                    'saldo': 0.0
                }

                # Revisar celda base del mes (start_c)
                v_base = sheet.cell(r, start_c).value
                n_base = clean_num(v_base)
                if n_base > 0:
                    month_fees.append(n_base)

                if has_banco_header:
                    # Formato estándar de 8 columnas
                    for c in range(start_c, end_c + 1):
                        h = str(sheet.cell(1, c).value or '').strip().upper()
                        v = sheet.cell(r, c).value
                        if v is None: continue

                        if 'FACTURA' in h:
                            m_info['factura'] = clean_val(v)
                        elif 'FECHA A PAGAR' in h:
                            m_info['fecha_a_pagar'] = clean_val(v)
                        elif 'FECHA DE PAGO' in h:
                            m_info['fecha_pago'] = clean_val(v)
                        elif h == 'PAGO':
                            m_info['pago'] = clean_num(v)
                        elif 'BANCO' in h:
                            m_info['banco'] = clean_val(v)
                        elif 'COD' in h:
                            m_info['cod'] = clean_val(v)
                        elif 'SALDO' in h:
                            m_info['saldo'] = clean_num(v)
                else:
                    # Formato heredado de 7 columnas (desplazado sin encabezado BANCO)
                    col2 = sheet.cell(r, start_c + 1).value # Factura o Fecha
                    col3 = sheet.cell(r, start_c + 2).value # Fecha
                    col4 = sheet.cell(r, start_c + 3).value # Pago
                    col5 = sheet.cell(r, start_c + 4).value # Banco
                    col6 = sheet.cell(r, start_c + 5).value # Cod
                    col7 = sheet.cell(r, start_c + 6).value # Saldo

                    if col2 is not None:
                        s_col2 = str(col2).strip()
                        if isinstance(col2, (datetime, date)) or (len(s_col2) >= 8 and ('/' in s_col2 or '-' in s_col2) and not s_col2.startswith('001') and not s_col2.startswith('C')):
                            m_info['fecha_a_pagar'] = clean_val(col2)
                        else:
                            m_info['factura'] = clean_val(col2)
                    if col3 is not None:
                        m_info['fecha_pago'] = clean_val(col3)
                    elif m_info['fecha_a_pagar'] and not m_info['factura']:
                        m_info['fecha_pago'] = m_info['fecha_a_pagar']
                        m_info['fecha_a_pagar'] = None

                    m_info['pago'] = clean_num(col4)
                    m_info['banco'] = clean_val(col5)
                    m_info['cod'] = clean_val(col6)
                    m_info['saldo'] = clean_num(col7)

                # Ajustes por coherencia de tipos
                if m_info['pago'] == 0.0 and isinstance(m_info['banco'], (int, float)):
                    m_info['pago'] = clean_num(m_info['banco'])
                    m_info['banco'] = None
                if m_info['banco'] is None and isinstance(m_info['pago'], str) and any(b in m_info['pago'].upper() for b in ['PICHINCHA', 'JEP', 'EFECTIVO']):
                    m_info['banco'] = m_info['pago']
                    m_info['pago'] = 0.0

                meses_data[m_name.lower()] = m_info

                if (m_info['pago'] > 0 or m_info['saldo'] > 0 or n_base > 0) and first_active_month_idx is None:
                    if m_name in months_order:
                        first_active_month_idx = months_order.index(m_name)

            # Determinar fecha_ingreso para cálculo de deudas
            if first_active_month_idx is not None:
                m_num = first_active_month_idx + 1
                fecha_ingreso_calc = f"2026-{m_num:02d}-01"
            else:
                fecha_ingreso_calc = "2026-01-01"

            # Buscar si el cliente ya existe
            cliente_existente = None
            if cod_raw and cod_raw.strip().upper() in by_cod:
                cliente_existente = by_cod[cod_raw.strip().upper()]
            elif nom_raw and _norm_str(nom_raw) in by_nombre:
                cliente_existente = by_nombre[_norm_str(nom_raw)]

            if cliente_existente:
                db_cliente = cliente_existente
                count_actualizados += 1
            else:
                db_cliente = models.ClienteExtra()
                db.add(db_cliente)
                count_nuevos += 1

            db_cliente.cod = cod_raw
            db_cliente.nombre_cliente = nom_raw
            db_cliente.contacto = contacto_raw
            db_cliente.proveedor = proveedor_raw
            db_cliente.usuario = usuario_raw
            db_cliente.contrasena = contrasena_raw
            db_cliente.cuentas = cuentas_raw
            db_cliente.mac_smart_one = mac_raw
            db_cliente.observaciones = obs_raw
            db_cliente.estado = estado_raw
            db_cliente.valor = valor_raw
            db_cliente.activo = activo_raw
            db_cliente.fecha_ingreso = fecha_ingreso_calc

            # Asignar meses a columnas del modelo (asignación exacta sin duplicación acumulativa)
            total_pagado_calc = 0.0
            for m_low, m_info in meses_data.items():
                setattr(db_cliente, f"{m_low}_factura", m_info['factura'])
                setattr(db_cliente, f"{m_low}_fecha_a_pagar", m_info['fecha_a_pagar'])
                setattr(db_cliente, f"{m_low}_fecha_pago", m_info['fecha_pago'])
                setattr(db_cliente, f"{m_low}_pago", m_info['pago'])
                setattr(db_cliente, f"{m_low}_banco", m_info['banco'])
                setattr(db_cliente, f"{m_low}_cod", m_info['cod'])
                setattr(db_cliente, f"{m_low}_saldo", m_info['saldo'])
                total_pagado_calc += m_info['pago']

            db_cliente.total_pagado = round(total_pagado_calc, 2)
            mes_actual_idx = datetime.now().month - 1
            mes_actual_str = months_order[mes_actual_idx].lower() if mes_actual_idx < len(months_order) else "octubre"
            oct_saldo = float(meses_data.get(mes_actual_str, {}).get('saldo', 0.0) or 0.0)
            is_eliminado = str(db_cliente.estado or "").strip().upper() == "ELIMINADO"
            if is_eliminado:
                db_cliente.saldo_pendiente = 0.0
                for m_name in months_order:
                    setattr(db_cliente, f"{m_name.lower()}_saldo", 0.0)
            else:
                db_cliente.saldo_pendiente = max(0.0, oct_saldo)

            # Flush para tener ID del cliente
            db.flush()

            # Sincronizar con PagoExtra para que el historial contenga cada pago
            for m_low, m_info in meses_data.items():
                if m_info['pago'] > 0:
                    f_pago_dt = None
                    if m_info['fecha_pago']:
                        try:
                            parts = m_info['fecha_pago'].split('/')
                            if len(parts) == 3:
                                f_pago_dt = datetime(int(parts[2]), int(parts[1]), int(parts[0]))
                        except Exception:
                            pass
                    if not f_pago_dt:
                        m_idx = months_order.index(m_low.upper()) + 1 if m_low.upper() in months_order else 1
                        f_pago_dt = datetime(2026, m_idx, 15)

                    pago_existente = db.query(models.PagoExtra).filter(
                        models.PagoExtra.cliente_id == db_cliente.id,
                        models.PagoExtra.mes_correspondiente == m_low.upper(),
                        models.PagoExtra.anulado == False
                    ).first()

                    if not pago_existente:
                        nuevo_pago = models.PagoExtra(
                            cliente_id=db_cliente.id,
                            monto=m_info['pago'],
                            metodo_pago=m_info['banco'] or "EFECTIVO",
                            mes_correspondiente=m_low.upper(),
                            referencia=m_info['cod'] or f"IMPORTACION {m_low.upper()}",
                            factura=m_info['factura'],
                            estado="Completado",
                            fecha_pago=f_pago_dt,
                            anulado=False
                        )
                        db.add(nuevo_pago)
                    else:
                        pago_existente.monto = m_info['pago']
                        if m_info['banco']: pago_existente.metodo_pago = m_info['banco']
                        if m_info['factura']: pago_existente.factura = m_info['factura']
                        pago_existente.fecha_pago = f_pago_dt

            # Actualizar diccionarios de lookup
            if cod_raw:
                by_cod[cod_raw.strip().upper()] = db_cliente
            if nom_raw:
                by_nombre[_norm_str(nom_raw)] = db_cliente

        db.commit()
        return {
            "message": f"Base de datos de Extras importada exitosamente: {count_nuevos} nuevos, {count_actualizados} actualizados.",
            "nuevos": count_nuevos,
            "actualizados": count_actualizados,
            "total": count_nuevos + count_actualizados
        }
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error al importar base de datos de extras: {str(e)}")


@router.get("/download-db", dependencies=[Depends(require_role(["administrador", "secretario"]))])
def download_database_extras(db: Session = Depends(get_db)):
    """
    Exporta la base de datos de Clientes Extras a un archivo Excel (.xlsx)
    con la estructura exacta de PLATAFORMA y los 12 meses.
    """
    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "PLATAFORMA"

        meses = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]

        # Encabezados
        headers = ["COD", "NOMBRE CLIENTE", "CONTACTO", "PROVEEDOR", "USUARIO", "CONTRASEÑA", "CUENTAS", "MAC SMART ONE", "OBSERVACIONES", "ESTADO", "VALOR", "ACTIVO"]
        for m in meses:
            headers.extend([m, "FACTURA", "FECHA A PAGAR", "FECHA DE PAGO", "PAGO", "BANCO", "COD", "SALDO"])

        ws.append(headers)

        # Estilo de encabezados
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        extras = db.query(models.ClienteExtra).all()
        # Ordenar por cod numérico si es posible
        def sort_key(e):
            if not e.cod: return 999999
            digits = re.findall(r'\d+', str(e.cod))
            return int(digits[0]) if digits else 999999
        extras.sort(key=sort_key)

        for e in extras:
            row = [
                e.cod,
                e.nombre_cliente,
                e.contacto,
                e.proveedor or "OPSATEL",
                e.usuario,
                e.contrasena,
                e.cuentas,
                e.mac_smart_one,
                e.observaciones,
                e.estado,
                float(e.valor or 0),
                e.activo or "SI"
            ]
            for m in meses:
                m_l = m.lower()
                row.extend([
                    float(e.valor or 0),
                    getattr(e, f"{m_l}_factura", None),
                    getattr(e, f"{m_l}_fecha_a_pagar", None),
                    getattr(e, f"{m_l}_fecha_pago", None),
                    float(getattr(e, f"{m_l}_pago", 0) or 0),
                    getattr(e, f"{m_l}_banco", None),
                    getattr(e, f"{m_l}_cod", None),
                    float(getattr(e, f"{m_l}_saldo", 0) or 0)
                ])
            ws.append(row)

        ws.freeze_panes = "C2"

        # Ajuste de anchos de columna
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 35)

        stream = io.BytesIO()
        wb.save(stream)
        stream.seek(0)

        filename = f"Base_Datos_Extras_General_{datetime.now().strftime('%Y%m%d')}.xlsx"
        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error al exportar base de datos de extras: {str(e)}")

